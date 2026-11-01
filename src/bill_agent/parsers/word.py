"""Word 解析（四.1：提取正文和表格）。

- .docx：python-docx，按文档顺序输出段落和表格
- .doc ：先用 LibreOffice 转成 .docx
- 页眉、页脚、文本框里的内容也收进 text（公司抬头经常在这些位置）
- 合并单元格保持重复值，保证列与表头对齐（与 Excel 处理方式一致）
"""

from __future__ import annotations

import io

from bill_agent.models import FileType, RawDocument, RawTable

from .base import InputFile, ParseError, ParseResult
from .convert import convert_with_soffice
from .detect import DetectedType, detect
from .table_utils import build_table, format_cell, grid_to_text

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _table_grid(table) -> list[list[str | None]]:
    return [[format_cell(cell.text) for cell in row.cells] for row in table.rows]


def _collect_tables(table, name: str, out: list[RawTable], texts: list[str]) -> None:
    grid = _table_grid(table)
    out.append(build_table(name, grid))
    texts.append(f"[{name}]\n{grid_to_text(grid)}")
    # 单元格内嵌套的表格
    k = 0
    for row in table.rows:
        for cell in row.cells:
            for inner in cell.tables:
                k += 1
                _collect_tables(inner, f"{name}.{k}", out, texts)


def _docx_to_document(data: bytes, f: InputFile, warnings: list[str]) -> RawDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ParseError(f"无法打开 docx: {e}") from e

    texts: list[str] = []
    tables: list[RawTable] = []

    # 页眉
    seen_hf: set[str] = set()
    for sec in d.sections:
        t = "\n".join(p.text for p in sec.header.paragraphs if p.text.strip())
        if t and t not in seen_hf:
            seen_hf.add(t)
            texts.append(f"## Header\n{t}")

    # 正文：按文档顺序
    n_table = 0
    for child in d.element.body.iterchildren():
        if child.tag == f"{_W}p":
            t = Paragraph(child, d).text.strip()
            if t:
                texts.append(t)
        elif child.tag == f"{_W}tbl":
            n_table += 1
            _collect_tables(Table(child, d), f"table-{n_table}", tables, texts)

    # 文本框
    box_texts = []
    for p in d.element.body.iter(f"{_W}txbxContent"):
        t = "".join(x.text or "" for x in p.iter(f"{_W}t")).strip()
        if t:
            box_texts.append(t)
    if box_texts:
        texts.append("## Text boxes\n" + "\n".join(dict.fromkeys(box_texts)))

    # 页脚
    for sec in d.sections:
        t = "\n".join(p.text for p in sec.footer.paragraphs if p.text.strip())
        if t and t not in seen_hf:
            seen_hf.add(t)
            texts.append(f"## Footer\n{t}")

    if not texts:
        warnings.append("文档没有文本内容")
    return RawDocument(
        source_id=f.source_id,
        source_path=f.source_path,
        parent_id=f.parent_id,
        file_type=FileType.WORD,
        text="\n".join(texts),
        tables=[t for t in tables if t.rows or t.header],
        warnings=warnings,
    )


class WordParser:
    def parse(self, f: InputFile, router) -> ParseResult:
        warnings: list[str] = []
        data = f.data
        if detect(f.data, f.name, f.content_type) is DetectedType.DOC:
            data = convert_with_soffice(f.data, "doc", "docx")
            if data is None:
                raise ParseError("无法解析 .doc：没有可用的 LibreOffice 进行转换")
            warnings.append("已通过 LibreOffice 将 .doc 转为 .docx")
        return ParseResult(documents=[_docx_to_document(data, f, warnings)])
