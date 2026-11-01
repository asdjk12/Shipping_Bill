"""Excel 解析（四.1：读取工作表、表头及费用明细行）。

- .xlsx：openpyxl（data_only=True 读公式的缓存结果）
- .xls ：xlrd；未安装时用 LibreOffice 转成 .xlsx
- 合并单元格展开为左上角的值，避免表头或费用名只出现在第一格
- 每个 sheet 一张 RawTable；全部 sheet 的内容同时拼成 text，表头以上的
  "Invoice No: xxx" 等信息也保留在 text 里
"""

from __future__ import annotations

import io

from bill_agent.models import FileType, RawDocument

from .base import InputFile, ParseError, ParseResult
from .convert import convert_with_soffice
from .detect import DetectedType, detect
from .table_utils import build_table, format_cell, grid_to_text

Grid = list[list[str | None]]


def _read_xlsx(data: bytes, warnings: list[str]) -> list[tuple[str, Grid]]:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    except Exception as e:
        raise ParseError(f"无法打开 xlsx: {e}") from e
    sheets: list[tuple[str, Grid]] = []
    formula_without_cache = False
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            warnings.append(f"跳过隐藏工作表: {ws.title}")
            continue
        grid: Grid = [[format_cell(c.value) for c in row] for row in ws.iter_rows()]
        for rng in ws.merged_cells.ranges:
            top = grid[rng.min_row - 1][rng.min_col - 1] if rng.min_row - 1 < len(grid) else None
            for r in range(rng.min_row - 1, min(rng.max_row, len(grid))):
                for c in range(rng.min_col - 1, min(rng.max_col, len(grid[r]))):
                    grid[r][c] = top
        sheets.append((ws.title, grid))
    # data_only 模式下，未经 Excel 保存过的公式单元格会是 None
    try:
        wb_f = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True)
        for ws_f, (title, grid) in zip((w for w in wb_f.worksheets if w.sheet_state == "visible"), sheets):
            for r_i, row in enumerate(ws_f.iter_rows()):
                for c_i, cell in enumerate(row):
                    v = getattr(cell, "value", None)
                    if isinstance(v, str) and v.startswith("=") and r_i < len(grid) and c_i < len(grid[r_i]) and grid[r_i][c_i] is None:
                        formula_without_cache = True
                        break
        wb_f.close()
    except Exception:
        pass
    if formula_without_cache:
        warnings.append("存在没有缓存结果的公式单元格，其值读为空")
    return sheets


def _read_xls(data: bytes, warnings: list[str]) -> list[tuple[str, Grid]]:
    try:
        import xlrd
    except ImportError:
        converted = convert_with_soffice(data, "xls", "xlsx")
        if converted is None:
            raise ParseError("无法解析 .xls：未安装 xlrd，且没有可用的 LibreOffice")
        warnings.append("已通过 LibreOffice 将 .xls 转为 .xlsx")
        return _read_xlsx(converted, warnings)

    try:
        book = xlrd.open_workbook(file_contents=data, formatting_info=True)
    except Exception:
        book = xlrd.open_workbook(file_contents=data)  # 部分文件不支持 formatting_info
    sheets: list[tuple[str, Grid]] = []
    for sh in book.sheets():
        if getattr(sh, "visibility", 0) != 0:
            warnings.append(f"跳过隐藏工作表: {sh.name}")
            continue
        grid: Grid = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    v = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    v = None
                row.append(format_cell(v))
            grid.append(row)
        for r0, r1, c0, c1 in getattr(sh, "merged_cells", []):
            top = grid[r0][c0]
            for r in range(r0, r1):
                for c in range(c0, c1):
                    grid[r][c] = top
        sheets.append((sh.name, grid))
    return sheets


class ExcelParser:
    def parse(self, f: InputFile, router) -> ParseResult:
        warnings: list[str] = []
        kind = detect(f.data, f.name, f.content_type)
        sheets = _read_xls(f.data, warnings) if kind is DetectedType.XLS else _read_xlsx(f.data, warnings)

        tables, text_parts = [], []
        for title, grid in sheets:
            t = build_table(title, grid)
            if not t.rows and not t.header:
                continue
            if t.header is None:
                warnings.append(f"工作表 {title} 未识别到费用表头")
            tables.append(t)
            text_parts.append(f"## Sheet: {title}\n{grid_to_text(grid)}")

        if not tables:
            warnings.append("工作簿中没有内容")
        doc = RawDocument(
            source_id=f.source_id,
            source_path=f.source_path,
            parent_id=f.parent_id,
            file_type=FileType.EXCEL,
            text="\n\n".join(text_parts),
            tables=tables,
            warnings=warnings,
        )
        return ParseResult(documents=[doc])
