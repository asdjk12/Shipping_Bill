"""解析层测试用的样例文件，全部在代码中生成。"""

from __future__ import annotations

import io
from email.message import EmailMessage
from pathlib import Path

import pytest

SAMPLES = Path(__file__).parents[1] / "fixtures" / "samples"
TINY_PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 200  # 仅用于类型识别的小图片


def make_text_pdf(lines: list[str]) -> bytes:
    """生成带文本层的最小 PDF（样例原件不入库，测试不依赖它）。"""
    content = "BT /F1 11 Tf 50 750 Td 14 TL " + " ".join(
        "(" + l.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ") '" for l in lines
    ) + " ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


DN_PDF = make_text_pdf([
    "DEBIT NOTE", "D/C No. : DC-108478", "D/C DATE 10-24-2026",
    "DESCRIPTION UNIT RATE QTY P/C DEBIT(+)",
    "HANDLING CHARGE - 65.00 1.00 C 65.00",
    "HANDLING CHARGE - 35.00 1.00 C 35.00",
])


def make_xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Debit Note"
    ws["A1"] = "ACME LOGISTICS DEBIT NOTE"
    ws.merge_cells("A1:E1")
    ws["A2"] = "Invoice No: INV-2026-001"
    ws.append([])
    ws.append(["Description", "Qty", "Rate", "Amount", "Currency"])
    ws.append(["Handling Fee", 1, 40.0, 40.0, "USD"])
    ws.append(["THC", 2, 125.5, 251.0, "USD"])
    ws.append(["Total", None, None, 291.0, "USD"])
    hidden = wb.create_sheet("secret")
    hidden["A1"] = "should not appear"
    hidden.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_docx() -> bytes:
    import docx

    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "BETA SHIPPING CO., LTD"
    d.sections[0].footer.paragraphs[0].text = "Bank details on request"
    d.add_paragraph("Debit Note No: DN-2026-088")
    d.add_paragraph("B/L No: BL123456")
    t = d.add_table(rows=3, cols=4)
    for r, row in enumerate([["Charge", "Amount", "Curr", "P/C"],
                             ["Doc Fee", "50.00", "USD", "PP"],
                             ["Storage Fee", "1,200.00", "CNY", "CC"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    d.add_paragraph("Please remit within 30 days.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_inner_eml() -> EmailMessage:
    inner = EmailMessage()
    inner["Subject"] = "原始账单"
    inner["From"] = "agent@beta.com"
    inner["To"] = "ops@pilot.com"
    inner.set_content("请见附件。")
    inner.add_attachment(make_docx(), maintype="application",
                         subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
                         filename="DN-2026-088.docx")
    return inner


def make_eml() -> bytes:
    msg = EmailMessage()
    msg["Subject"] = "=?utf-8?b?6LSm5Y2V5pu05pawIERDLTEwODQ3OA==?="  # "账单更新 DC-108478"
    msg["From"] = "Neptune LAX <billing@neptune.com>"
    msg["To"] = "Pilot Ops <ops@pilot.com>, finance@pilot.com"
    msg["Cc"] = "boss@pilot.com"
    msg["Date"] = "Wed, 28 Oct 2026 10:00:00 +0800"
    msg["Message-ID"] = "<abc@neptune.com>"
    msg.set_content("Hi,\n请将 Handling Fee 改为 USD 40。\n附件为最新账单。")
    msg.add_alternative(
        "<html><body><p>Hi,</p><p>请将 Handling Fee 改为 USD 40。</p>"
        "<table><tr><th>Charge</th><th>Amount</th><th>Currency</th></tr>"
        "<tr><td>Handling Fee</td><td>40.00</td><td>USD</td></tr></table>"
        '<img src="cid:logo"></body></html>',
        subtype="html",
    )
    msg.get_body(("html",)).add_related(TINY_PNG, maintype="image", subtype="png", cid="<logo>")
    xlsx = make_xlsx()
    msg.add_attachment(xlsx, maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename="charges.xlsx")
    msg.add_attachment(xlsx, maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename="charges_copy.xlsx")
    # 扩展名错误：实际是 PDF
    msg.add_attachment(DN_PDF, maintype="application",
                       subtype="octet-stream", filename="DN_DC-108478.dat")
    msg.add_attachment(b"just notes", maintype="text", subtype="plain", filename="notes.txt")
    msg.add_attachment(make_inner_eml())  # 转发的邮件 message/rfc822
    return msg.as_bytes()


@pytest.fixture
def xlsx_bytes() -> bytes:
    return make_xlsx()


@pytest.fixture
def docx_bytes() -> bytes:
    return make_docx()


@pytest.fixture
def eml_bytes() -> bytes:
    return make_eml()
